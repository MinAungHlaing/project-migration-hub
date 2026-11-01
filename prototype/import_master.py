# Import products + Actual Base from Ko's corrected master file, and areas from the customer template.
# Usage: python import_master.py [master.xlsx] [template.xlsx]
import os, re, sqlite3, sys
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
DB = os.environ.get('GTM_DB', os.path.join(BASE, 'data', 'gtm.db'))
IMG = os.path.join(BASE, 'static', 'product-images')
MASTER = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser('~/workspace/user/files/sale_price_catalog_confidential_test.xlsx')
TEMPLATE = sys.argv[2] if len(sys.argv) > 2 else os.path.expanduser('~/workspace/user/files/GTM_CUSTOMER_TEMPLATE.xlsx')

sys.path.insert(0, BASE)
os.environ['GTM_DB'] = DB
import app as appmod
appmod.init_db()

con = sqlite3.connect(DB)
if con.execute("SELECT COUNT(*) FROM products").fetchone()[0] > 0:
    print('products already imported - skipping'); sys.exit(0)

wb = openpyxl.load_workbook(MASTER, data_only=True)
ws = wb.worksheets[0]
actual = {}
if len(wb.worksheets) > 1:
    for r in wb.worksheets[1].iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            try: actual[str(r[0]).strip()] = float(r[1]) if r[1] is not None else None
            except (TypeError, ValueError): actual[str(r[0]).strip()] = None

def num(v):
    try: return float(v) if v is not None and v != '' else 0
    except (TypeError, ValueError): return 0

n = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[1]: continue
    sku = re.sub(r'\s+', '', str(r[1]))
    name = str(r[2]).strip() if r[2] is not None else ''
    img = f'{sku}.webp' if os.path.exists(os.path.join(IMG, f'{sku}.webp')) else None
    con.execute("""INSERT INTO products(sku,category,name,unit,upc,supplier,base_price,actual_base,retail,wholesale,b2c,image,legacy_status,sellable)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (sku, r[0], name, r[6], num(r[5]) or None, r[13], num(r[9]),
                 actual.get(name), num(r[3]), num(r[7]), num(r[8]), img, r[4],
                 0 if str(r[4]).strip() == 'Out Of Stock' else 1))
    pid = con.execute("SELECT last_insert_rowid()").fetchone()[0]
    con.execute("INSERT OR IGNORE INTO stock(product_id,qty) VALUES(?,0)", (pid,))
    n += 1

con.execute("INSERT OR REPLACE INTO counters(name,val) VALUES('product_seq',360)")
con.execute("INSERT OR IGNORE INTO counters(name,val) VALUES('order_seq',0)")
con.execute("INSERT OR IGNORE INTO counters(name,val) VALUES('purchase_seq',0)")

areas = 0
if os.path.exists(TEMPLATE):
    twb = openpyxl.load_workbook(TEMPLATE, data_only=True)
    if 'Data' in twb.sheetnames:
        for r in twb['Data'].iter_rows(min_row=2, values_only=True):
            if r[0]:
                con.execute("INSERT OR IGNORE INTO areas(name) VALUES(?)", (str(r[0]).strip(),)); areas += 1
con.commit()
print(f'imported {n} products, {areas} areas')
print('with images:', con.execute("SELECT COUNT(*) FROM products WHERE image IS NOT NULL").fetchone()[0])
print('with actual base:', con.execute("SELECT COUNT(*) FROM products WHERE actual_base IS NOT NULL").fetchone()[0])
con.close()
